"""Rotas HTTP - PGM Porto Velho - Subprocuradoria Contenciosa."""
from flask import Blueprint, render_template, request, jsonify, current_app
from functools import wraps
import os, json, requests as http
from datetime import datetime, date, timedelta

bp = Blueprint('main', __name__)

# Configurações do Supabase
SUPABASE_URL = os.environ.get('SUPABASE_URL', 'https://vmbzykywtgzyxmoxogel.supabase.co')
SUPABASE_KEY = os.environ.get('SUPABASE_KEY', 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InZtYnp5a3l3dGd6eXhtb3hvZ2VsIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc3NjM3NjAyMywiZXhwIjoyMDkxOTUyMDIzfQ.fJlvsVIPLkIg5IXNx1uYfqa5pDj1B8DbRQIUiiTpcEo')

def _sb_headers():
    return {'apikey': SUPABASE_KEY, 'Authorization': f'Bearer {SUPABASE_KEY}', 'Content-Type': 'application/json', 'Prefer': 'return=representation'}

def _sb_get(table, params=''):
    try:
        r = http.get(f'{SUPABASE_URL}/rest/v1/{table}?{params}', headers=_sb_headers(), timeout=10)
        return r.json() if r.ok else []
    except: return []

def cache_get(chave):
    try:
        rows = _sb_get('dados_cache', f'chave=eq.{chave}&select=valor')
        return rows[0].get('valor') if rows else None
    except: return None

def cache_set(chave, valor):
    try:
        h = _sb_headers()
        h['Prefer'] = 'resolution=merge-duplicates,return=minimal'
        http.post(f'{SUPABASE_URL}/rest/v1/dados_cache', headers=h, json={'chave': chave, 'valor': valor, 'atualizado': datetime.utcnow().isoformat()}, timeout=10)
    except: pass

# --- LÓGICA DE PROCESSAMENTO DINÂMICO ---

def _recalcular_prazos(lista_completa):
    """Recalcula o status de cada processo comparando com a data de hoje."""
    hoje = date.today()
    vencidos, proximos, cumpridos = [], [], []
    perf = {}
    manuais = cache_get('cumpridos_manuais') or []

    for p in lista_completa:
        try:
            dt_prazo = date.fromisoformat(p['data_iso'])
        except: continue
        
        proc = p['processo']
        resp = p['responsavel']
        ja_cumprido = p.get('ja_cumprido', False) or proc in manuais
        
        if resp not in perf: perf[resp] = {'total':0, 'cumpridos':0, 'criticos':0}
        perf[resp]['total'] += 1

        if ja_cumprido:
            perf[resp]['cumpridos'] += 1
            cumpridos.append(p)
            continue

        diff = (dt_prazo - hoje).days
        p['dias'] = abs(diff)
        
        if diff < 0:
            perf[resp]['criticos'] += 1
            vencidos.append(p)
        elif diff <= 7:
            proximos.append(p)

    # Formatação para o Dashboard
    perf_list = []
    for r, d in perf.items():
        t, c = d['total'], d['cumpridos']
        perf_list.append({'responsavel': r, 'total': t, 'cumpridos': c, 'taxa': round(c/t*100, 1) if t > 0 else 0, 'criticos': d['criticos']})

    return {
        'stats': {'total': len(lista_completa), 'vencidos': len(vencidos), 'proximos': len(proximos), 'cumpridos': len(cumpridos), 'taxa': round(len(cumpridos)/len(lista_completa)*100, 1) if lista_completa else 0, 'ultima_atualizacao': hoje.strftime('%d/%m/%Y')},
        'performance': sorted(perf_list, key=lambda x: x['taxa'], reverse=True),
        'proximos': sorted(proximos, key=lambda x: x['dias']),
        'vencidos': sorted(vencidos, key=lambda x: x['dias'], reverse=True)
    }

def _parse_xlsx(file_obj):
    import openpyxl, warnings
    from datetime import datetime, date
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    
    # Tenta achar a aba correta (Prazos 2026 ou Pauta)
    ws = None
    for name in ['Prazos 2026', 'Pauta', 'Planilha1']:
        if name in wb.sheetnames: ws = wb[name]; break
    if not ws: ws = wb.active

    lista_completa = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[1] or not row[4]: continue
        prazo_raw = row[1]
        if isinstance(prazo_raw, datetime): prazo_d = prazo_raw.date()
        elif isinstance(prazo_raw, date): prazo_d = prazo_raw
        else: continue
        
        lista_completa.append({
            'processo': str(row[4]).strip(),
            'parte': str(row[5]).strip()[:60] if row[5] else '',
            'responsavel': str(row[2]).strip().upper() if row[2] else 'SEM RESPONSAVEL',
            'prazo': prazo_d.strftime('%d/%m/%Y'),
            'data_iso': prazo_d.isoformat(),
            'vara': str(row[6]).strip() if row[6] else '',
            'ja_cumprido': str(row[13]).strip().upper() in ('SIM', 'PARCIAL', 'PREJUDICADO')
        })
    return lista_completa

# --- ROTAS ---

@bp.route('/api/upload', methods=['POST'])
def upload_file():
    file = request.files.get('file')
    if not file: return jsonify({'error': 'Arquivo obrigatorio'}), 400
    try:
        lista = _parse_xlsx(file)
        cache_set('lista_mestra', lista)
        res = _recalcular_prazos(lista)
        cache_set('stats', res['stats'])
        cache_set('filename', file.filename)
        return jsonify({'success': True, 'stats': res['stats']})
    except Exception as e: return jsonify({'error': str(e)}), 500

@bp.route('/api/dashboard')
def get_dashboard():
    lista = cache_get('lista_mestra')
    if not lista: return jsonify({'sem_dados': True})
    return jsonify(_recalcular_prazos(lista))

@bp.route('/api/criticos')
def get_criticos():
    lista = cache_get('lista_mestra')
    if not lista: return jsonify({'sem_dados': True})
    res = _recalcular_prazos(lista)
    return jsonify({'vencidos': res['vencidos'], 'proximos': res['proximos']})

@bp.route('/painel')
def painel(): return render_template('dashboard.html')

@bp.route('/')
def index(): return render_template('dashboard.html')
